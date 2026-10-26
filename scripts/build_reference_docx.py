"""Builds tools/reference.docx: the Word style template pandoc uses when
compiling the report, encoding the brief's format requirements (Arial 12,
double-spaced body text, numbered pages) so every build is consistent.

Usage: .venv/bin/python scripts/build_reference_docx.py <base-docx> <out-docx>
"""
import sys

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Pt


def set_font(style, name, size_pt, bold=None):
    style.font.name = name
    style.font.size = Pt(size_pt)
    if bold is not None:
        style.font.bold = bold
    # Also set the East Asian font so Word doesn't fall back to Calibri.
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), name)


def set_double_spacing(style):
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE


def add_page_number_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.alignment = 1  # center

    run = paragraph.add_run()
    fld_char_begin = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    instr_text = run._r.makeelement(qn("w:instrText"), {})
    instr_text.text = "PAGE"
    fld_char_separate = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "separate"})
    fld_char_end = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})

    run._r.append(fld_char_begin)
    run._r.append(instr_text)
    run._r.append(fld_char_separate)
    run._r.append(fld_char_end)


def add_table_borders(style):
    # Pandoc's "Table" style has no cell borders, so tables print as loose
    # text in Word. Give every cell a thin single line.
    tbl_pr = style.element.find(qn("w:tblPr"))
    borders = tbl_pr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(borders.makeelement(qn(f"w:{edge}"), {
            qn("w:val"): "single", qn("w:sz"): "4",
            qn("w:space"): "0", qn("w:color"): "000000",
        }))
    # Schema order: tblBorders goes before tblCellMar.
    cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if cell_mar is not None:
        cell_mar.addprevious(borders)
    else:
        tbl_pr.append(borders)


def get_style(doc, name):
    # Styles.__contains__ / __getitem__ unreliable across style name casing
    # quirks in pandoc's default template, so match by exact .name instead.
    for style in doc.styles:
        if style.name == name:
            return style
    return None


def main():
    base_path, out_path = sys.argv[1], sys.argv[2]
    doc = Document(base_path)

    body_styles = ["Normal", "Body Text", "First Paragraph", "Compact"]
    for name in body_styles:
        style = get_style(doc, name)
        if style is not None:
            set_font(style, "Arial", 12)
            set_double_spacing(style)

    heading_sizes = {"Heading 1": 16, "Heading 2": 14, "Heading 3": 12, "Title": 20}
    for name, size in heading_sizes.items():
        style = get_style(doc, name)
        if style is not None:
            set_font(style, "Arial", size, bold=True)
            # Headings stay single-spaced even though body text is double.
            style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    # Each of the 8 report sections starts with a Heading 1 - start it on a
    # fresh page so sections don't run together in the compiled report.
    heading1 = get_style(doc, "Heading 1")
    if heading1 is not None:
        heading1.paragraph_format.page_break_before = True

    table = get_style(doc, "Table")
    if table is not None:
        add_table_borders(table)

    add_page_number_footer(doc)

    doc.save(out_path)
    print(f"Wrote {out_path}")


if __name__ == "__main__":
    main()
